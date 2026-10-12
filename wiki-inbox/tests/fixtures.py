"""Generate binary fixtures at test time (keeps the repo free of binaries)."""

from pathlib import Path


def make_pdf(path: Path, text: str) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(72, 750, text)
    c.showPage()
    c.save()
    return path


def make_docx(path: Path, heading: str, body: str) -> Path:
    import docx

    d = docx.Document()
    d.add_heading(heading, level=1)
    d.add_paragraph(body)
    d.save(str(path))
    return path


HTML = """<!doctype html><html><head><title>Fixture page</title></head>
<body><h1>Gradient descent</h1><p>An <b>optimisation</b> algorithm.</p>
<ul><li>step</li><li>repeat</li></ul></body></html>"""
